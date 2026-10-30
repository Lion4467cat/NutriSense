"""Build an identity-grouped, label-only manifest for FinnFoodPics 2026.

The output references local image files; it never copies source images. The
calorie tier is a target label and must not be used as a model input. The
derived mass is nominal (scaled from the workbook's 120 kcal portion), not a
per-image scale measurement and must remain separate from WED measured grams.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import re
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
RELEASE = ROOT / "datasets" / "raw" / "portion" / "finnfoodpics" / "versions" / "2026"
IMAGE_ROOT = RELEASE / "1_Image database"
WORKBOOK = RELEASE / "3_FinnFoodPics _info_database_2026.xlsx"
DEFAULT_OUT = ROOT / "datasets" / "processed" / "finnfoodpics_v2026_manifest.csv"

IMAGE_NAME = re.compile(
    r"(?P<index>\d+)_(?P<category>Carbohydrate|Fat|Combo)_"
    r"(?P<kcal>40|80|120|160|200|240|280|320|360|400)\.jpg",
    re.IGNORECASE,
)
SPLIT_SEED = "finnfoodpics-2026-identity-split-v1"


def _norm_category(value: object) -> str:
    return str(value).strip().casefold()


def _workbook_rows() -> dict[tuple[str, int], dict[str, object]]:
    if not WORKBOOK.is_file():
        raise FileNotFoundError(f"FinnFoodPics workbook not found: {WORKBOOK}")
    workbook = load_workbook(WORKBOOK, read_only=True, data_only=True)
    if "120 kcal_CLEAN" not in workbook.sheetnames:
        raise ValueError("2026 workbook is missing the '120 kcal_CLEAN' sheet")
    sheet = workbook["120 kcal_CLEAN"]
    header = [str(cell.value).strip() if cell.value is not None else ""
              for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    index = {name: pos for pos, name in enumerate(header) if name}
    required = {"Category", "Index", "Complete_Description",
                "Energy_density_(kcal/g)", "Portion_(g/120 kcal)"}
    missing = required - index.keys()
    if missing:
        raise ValueError(f"FinnFoodPics workbook missing columns: {sorted(missing)}")

    rows: dict[tuple[str, int], dict[str, object]] = {}
    for values in sheet.iter_rows(min_row=2, values_only=True):
        category = values[index["Category"]]
        food_index = values[index["Index"]]
        if category is None or food_index is None:
            continue
        key = (_norm_category(category), int(food_index))
        if key in rows:
            raise ValueError(f"Duplicate FinnFoodPics workbook identity: {key}")
        rows[key] = {
            "description": values[index["Complete_Description"]],
            "energy_density": values[index["Energy_density_(kcal/g)"]],
            "portion_g_120kcal": values[index["Portion_(g/120 kcal)"]],
        }
    workbook.close()
    return rows


def _identity_split(group_id: str, category_groups: list[str]) -> str:
    """Stable category-stratified 70/15/15 split at snack identity level."""
    ordered = sorted(
        category_groups,
        key=lambda value: hashlib.sha256(f"{SPLIT_SEED}:{value}".encode()).hexdigest(),
    )
    # Keep at least one held-out identity for each category while preserving
    # every tier from an identity in the same partition.
    n_test = max(1, round(len(ordered) * 0.15))
    n_validation = max(1, round(len(ordered) * 0.15))
    position = ordered.index(group_id)
    if position < len(ordered) - n_test - n_validation:
        return "train"
    if position < len(ordered) - n_test:
        return "validation"
    return "test"


def build_manifest() -> list[dict[str, object]]:
    if not IMAGE_ROOT.is_dir():
        raise FileNotFoundError(f"FinnFoodPics 2026 images not found: {IMAGE_ROOT}")
    workbook = _workbook_rows()
    images = sorted(IMAGE_ROOT.rglob("*.jpg"))
    if len(images) != 720:
        raise ValueError(f"Expected 720 FinnFoodPics 2026 images; found {len(images)}")

    parsed: list[dict[str, object]] = []
    groups_by_category: dict[str, set[str]] = defaultdict(set)
    group_keys: set[tuple[str, int]] = set()
    image_counts: Counter[tuple[str, int]] = Counter()
    tier_sets: dict[tuple[str, int], set[int]] = defaultdict(set)
    for path in images:
        match = IMAGE_NAME.fullmatch(path.name)
        if match is None:
            raise ValueError(f"Unrecognized FinnFoodPics image filename: {path.name}")
        category = match.group("category").title()
        food_index = int(match.group("index"))
        kcal = int(match.group("kcal"))
        key = (_norm_category(category), food_index)
        group_id = f"{category.casefold()}_{food_index}"
        labels = workbook.get(key)
        groups_by_category[category].add(group_id)
        group_keys.add(key)
        image_counts[key] += 1
        tier_sets[key].add(kcal)
        parsed.append({
            "observation_id": f"{group_id}_{kcal}kcal",
            "image_path": path.relative_to(ROOT).as_posix(),
            "source_release": "FinnFoodPics-2026",
            "license_ref": (
                "datasets/raw/portion/finnfoodpics/versions/2026/"
                "2_FinnFoodPics license 2026.pdf"
            ),
            "food_identity": group_id,
            "source_category": category,
            "description": labels["description"] if labels else None,
            "portion_kcal_label": kcal,
            "energy_density_kcal_per_g": labels["energy_density"] if labels else None,
            "portion_g_at_120kcal": labels["portion_g_120kcal"] if labels else None,
            "nominal_mass_g_from_120kcal": (
                float(labels["portion_g_120kcal"]) * kcal / 120
                if labels and labels["portion_g_120kcal"] is not None else None
            ),
            "workbook_labels_available": labels is not None,
            "mass_target_kind": (
                "nominal_scaled_from_120kcal_workbook_portion"
                if labels else "unavailable_calorie_tier_only"
            ),
            "group_id": group_id,
        })

    expected_tiers = {40, 80, 120, 160, 200, 240, 280, 320, 360, 400}
    for key in group_keys:
        if image_counts[key] != 10 or tier_sets[key] != expected_tiers:
            raise ValueError(
                f"Identity {key} must have one image at each of ten tiers; "
                f"count={image_counts[key]} tiers={sorted(tier_sets[key])}"
            )
    for key in workbook:
        if key not in group_keys:
            raise ValueError(f"Workbook identity has no images in 2026 release: {key}")

    groups_by_category_sorted = {
        category: sorted(groups) for category, groups in groups_by_category.items()
    }
    for row in parsed:
        category = str(row["source_category"])
        row["split"] = _identity_split(
            str(row["group_id"]), groups_by_category_sorted[category]
        )
    return parsed


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", type=Path, default=DEFAULT_OUT)
    args = parser.parse_args(argv)
    rows = build_manifest()
    args.out.parent.mkdir(parents=True, exist_ok=True)
    columns = list(rows[0])
    with args.out.open("w", newline="", encoding="utf-8") as stream:
        writer = csv.DictWriter(stream, fieldnames=columns)
        writer.writeheader()
        writer.writerows(rows)
    image_counts = Counter(str(row["split"]) for row in rows)
    missing_rows = [row for row in rows if not row["workbook_labels_available"]]
    missing_workbook = len(missing_rows)
    missing_groups = sorted({str(row["food_identity"]) for row in missing_rows})
    identity_counts = Counter(
        (str(row["source_category"]), str(row["split"]), str(row["group_id"]))
        for row in rows
    )
    category_identities: Counter[tuple[str, str]] = Counter(
        (category, split)
        for category, split, _ in identity_counts
    )
    print(f"rows={len(rows)} identities={len(identity_counts)} "
          f"images_without_workbook_labels={missing_workbook} output={args.out}")
    if missing_groups:
        print(f"identities_without_workbook_labels={','.join(missing_groups)}")
    for split in ("train", "validation", "test"):
        n_ids = sum(count for (_, group_split), count in category_identities.items()
                    if group_split == split)
        print(f"split={split} identities={n_ids} images={image_counts[split]}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
