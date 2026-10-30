"""Audit LeFood-Set V1 image/workbook correspondence and write a fixed split."""
from __future__ import annotations

import argparse
import csv
import hashlib
from collections import Counter
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
ARCHIVE = ROOT / "datasets/raw/portion/lefood-set/LeFood-Set_v1.zip"
WORKBOOK = ROOT / "datasets/raw/portion/lefood-set/metadata/data_original.xlsx"
IMAGE_DIR = ROOT / "datasets/raw/portion/lefood-set/images"
MANIFEST = Path(__file__).resolve().parent / "manifest.csv"
AUDIT = Path(__file__).resolve().parent / "DATA_AUDIT.md"
SPLIT_SEED = "nutrisense-lefood-category-split-20261030:"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--archive", type=Path, default=ARCHIVE)
    parser.add_argument("--workbook", type=Path, default=WORKBOOK)
    parser.add_argument("--manifest", type=Path, default=MANIFEST)
    parser.add_argument("--audit", type=Path, default=AUDIT)
    parser.add_argument("--extract", action="store_true",
                        help="extract valid before/after images into ignored raw storage")
    args = parser.parse_args()

    with ZipFile(args.archive) as archive:
        members = {Path(name).name: name for name in archive.namelist()
                   if name.lower().endswith((".jpg", ".jpeg", ".png"))}
        workbook = load_workbook(args.workbook, data_only=True, read_only=True)
        sheet = workbook.active
        headers = [str(value).strip() for value in next(
            sheet.iter_rows(min_row=1, max_row=1, values_only=True))]
        index = {name: position for position, name in enumerate(headers)}
        rows = list(sheet.iter_rows(min_row=2, values_only=True))

        valid = []
        missing_before = missing_after = invalid_target = 0
        for row in rows:
            before = str(row[index["Image Before Eaten"]]).strip()
            after = str(row[index["Image After Eaten"]]).strip()
            if before not in members:
                missing_before += 1
            if after not in members:
                missing_after += 1
            if before not in members or after not in members:
                continue
            before_g = row[index["Weight Before Eaten (g)"]]
            after_g = row[index["Weight After Eaten (g)"]]
            if not isinstance(before_g, (int, float)) or not isinstance(after_g, (int, float)):
                invalid_target += 1
                continue
            valid.append({
                "row_id": str(row[index["ID"]]),
                "food_name": str(row[index["Name of the food"]]).strip(),
                "before_name": before,
                "after_name": after,
                "before_g": float(before_g),
                "after_g": float(after_g),
                "observer_level": str(row[index["Visual Estimation by Observer (1-7)"]]),
            })

        food_names = sorted({row["food_name"] for row in valid}, key=lambda name: (
            hashlib.sha256(f"{SPLIT_SEED}{name}".encode("utf-8")).hexdigest(), name))
        n_train = int(len(food_names) * 0.70)
        n_dev = int(len(food_names) * 0.15)
        split_by_food = {
            name: ("train" if i < n_train else "dev" if i < n_train + n_dev else "test")
            for i, name in enumerate(food_names)
        }
        counts = Counter(split_by_food[row["food_name"]] for row in valid)
        categories = {split: [name for name in food_names if split_by_food[name] == split]
                      for split in ("train", "dev", "test")}

        if args.extract:
            for row in valid:
                for side, name in (("before", row["before_name"]),
                                   ("after", row["after_name"])):
                    out = IMAGE_DIR / side / f"{row['row_id']}_{Path(name).name}"
                    if not out.is_file():
                        out.parent.mkdir(parents=True, exist_ok=True)
                        out.write_bytes(archive.read(members[name]))
                    row[f"{side}_path"] = out.relative_to(ROOT).as_posix()

    args.manifest.parent.mkdir(parents=True, exist_ok=True)
    fields = ["row_id", "food_name", "before_path", "after_path", "before_g",
              "after_g", "observer_level", "group_id", "split", "target_kind",
              "target_unit", "license_ref"]
    with args.manifest.open("w", newline="", encoding="utf-8") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields)
        writer.writeheader()
        for row in valid:
            writer.writerow({
                "row_id": row["row_id"],
                "food_name": row["food_name"],
                "before_path": row.get("before_path", row["before_name"]),
                "after_path": row.get("after_path", row["after_name"]),
                "before_g": row["before_g"],
                "after_g": row["after_g"],
                "observer_level": row["observer_level"],
                "group_id": row["food_name"],
                "split": split_by_food[row["food_name"]],
                "target_kind": "source_reported_food_mass_before_after",
                "target_unit": "g",
                "license_ref": "LeFood-Set V1 CC BY 4.0; see docs/licenses.md",
            })

    lines = [
        "# LeFood-Set V1 data audit",
        "",
        f"Archive SHA-256: `{sha256(args.archive)}`.",
        f"Workbook SHA-256: `{sha256(args.workbook)}`.",
        f"Workbook data rows: {len(rows)}; complete before/after image pairs present in archive: {len(valid)}.",
        f"Rows with missing before filenames: {missing_before}; missing after filenames: {missing_after}; "
        f"complete rows with nonnumeric targets: {invalid_target}.",
        f"Food categories in complete pairs: {len(food_names)}.",
        "",
        "## Fixed category-group split",
        "",
        f"Seed: `{SPLIT_SEED}`; categories are assigned as whole groups.",
        "| Split | Food categories | Image pairs |",
        "|---|---:|---:|",
    ]
    for split in ("train", "dev", "test"):
        lines.append(f"| {split} | {len(categories[split])} | {counts[split]} |")
    for split in ("train", "dev", "test"):
        lines.extend(["", f"### {split} categories", "", ", ".join(categories[split])])
    lines.extend([
        "",
        "## Target semantics note",
        "",
        "The Mendeley record describes the before/after values as food weights in grams. "
        "The paper says scale measurements were made by weighing the plate containing food, "
        "but does not describe tare handling. Treat the source's food-weight labels as stated "
        "for exploratory scoring, and retain tare correction as an unresolved provenance caveat.",
        "",
    ])
    args.audit.write_text("\n".join(lines), encoding="utf-8")
    print(f"workbook_rows={len(rows)} complete_pairs={len(valid)} "
          f"missing_before={missing_before} missing_after={missing_after} "
          f"invalid_targets={invalid_target} categories={len(food_names)} split_pairs={dict(counts)}")
    print(f"manifest={args.manifest} sha256={sha256(args.manifest)}")
    print(f"audit={args.audit}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
